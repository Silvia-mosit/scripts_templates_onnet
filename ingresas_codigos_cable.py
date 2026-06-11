import openpyxl
from openpyxl.utils import column_index_from_string

ENTEL_CABLE_NAMES = {
    1001: "ENTEL ADSS 192",
    1002: "ENTEL ADSS 32",
    1003: "ENTEL ADSS 48 (4M x 12F)",
    1004: "ENTEL ADSS 48 (6M x 8F)",
    1005: "ENTEL ADSS 96",
    1006: "ENTEL 12",
    1007: "ENTEL 16 FO Indoor-Outdoor",
    1008: "ENTEL 192",
    1009: "ENTEL 2 (RISER)",
    1010: "ENTEL 24",
    1011: "ENTEL 24 (RISER)",
    1012: "ENTEL 24 SOPLADO (3M x 8F)",
    1013: "ENTEL 288",
    1014: "ENTEL 32 (4M x 8F)",
    1015: "ENTEL 32 RISER (4M x 8F)",
    1016: "ENTEL 48",
    1017: "ENTEL 48 (6M x 8F)",
    1018: "ENTEL 48 RISER (6M x 8F)",
    1019: "ENTEL 6",
    1020: "ENTEL 64 RISER (8M x 8F)",
    1021: "ENTEL 8 FO Indoor-Outdoor",
    1022: "ENTEL 96",
}


def _col_to_idx(ws, col: str) -> int:
    """Devuelve índice 1-based desde nombre de encabezado o letra Excel (A, B, ..., Z, AA, ...)."""
    for cell in ws[1]:
        if str(cell.value) == col:
            return cell.column
    try:
        return column_index_from_string(col.upper())
    except Exception:
        raise ValueError(f"Columna '{col}' no encontrada como encabezado ni como letra Excel.")


def map_cable_names(input_path: str, key_col: str, value_col: str, output_path: str = None):
    """
    Lee un xlsx con encabezado, mapea los valores de key_col a nombres usando
    ENTEL_CABLE_NAMES y los escribe en value_col, preservando el formato original.

    key_col y value_col aceptan nombre de encabezado ("codigo") o letra Excel ("M").
    """
    wb = openpyxl.load_workbook(input_path)
    ws = wb.active

    key_idx = _col_to_idx(ws, key_col)
    val_idx = _col_to_idx(ws, value_col)

    not_found = []
    for row in ws.iter_rows(min_row=2):
        key_cell = row[key_idx - 1]
        val_cell = row[val_idx - 1]

        try:
            key = int(key_cell.value)
        except (TypeError, ValueError):
            continue  # Saltar celdas vacías o de texto

        name = ENTEL_CABLE_NAMES.get(key)
        if name:
            val_cell.value = name
        else:
            val_cell.value = "CÓDIGO NO ENCONTRADO"
            not_found.append(key)

    if not_found:
        unique_not_found = list(dict.fromkeys(not_found))
        print(f"ADVERTENCIA: {len(unique_not_found)} código(s) no encontrado(s) en el diccionario: {unique_not_found}")

    out = output_path or input_path
    wb.save(out)
    print(f"Archivo guardado en: {out}")


if __name__ == "__main__":
    map_cable_names(
        input_path="data_xlsx/oriental_t11.xlsx",
        key_col="M",
        value_col="O",
    )
